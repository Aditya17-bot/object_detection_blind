# Build the Review-2 deck for BlindAssist in the VIT template supplied as
# "4 Review 2 PPT Template (1).pdf".
#
# The template is deliberately plain: white ground, black Times New Roman,
# one blue table on the literature-review slide, a grey footer carrying
# date / school / slide number. Nothing here changes that.

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "ppt_assets")
IMG = os.path.join(ASSETS, "img")
LOGO = os.path.join(ASSETS, "vit_logo.png")
OUT = os.environ.get("BLINDASSIST_PPT_OUT",
                     r"C:\adi\object_detection_blind\BlindAssist_Review2.pptx")

SERIF = "Times New Roman"
SANS = "Calibri"
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x7F, 0x7F, 0x7F)
TBL_HDR = RGBColor(0x44, 0x72, 0xC4)
TBL_A = RGBColor(0xD0, 0xD8, 0xE8)
TBL_B = RGBColor(0xE9, 0xED, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

DATE = "07-09-2026"
SCHOOL = "SCOPE"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def style(run, size=20, bold=False, italic=False, font=SERIF, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    # East-Asian / complex-script font names, so PowerPoint does not
    # substitute a different face for the same characters.
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def bullet(paragraph, level=0):
    """Give a paragraph a real round bullet at the requested indent level."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("indent", str(-228600))
    pPr.set("marL", str(228600 + level * 342900))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": "\u2022"})
    pPr.append(buFont)
    pPr.append(buChar)


def nobullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if pPr.find(qn("a:buNone")) is None:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def footer(slide, number):
    tf = textbox(slide, 1.0, 7.05, 3.0, 0.3)
    style(tf.paragraphs[0].add_run(), 11, font=SANS, color=GREY)
    tf.paragraphs[0].runs[0].text = DATE
    nobullet(tf.paragraphs[0])

    tf = textbox(slide, 5.2, 7.05, 3.0, 0.3)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = SCHOOL
    style(r, 11, font=SANS, color=GREY)
    nobullet(tf.paragraphs[0])

    tf = textbox(slide, 11.0, 7.05, 1.4, 0.3)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    r = tf.paragraphs[0].add_run()
    r.text = str(number)
    style(r, 11, font=SANS, color=GREY)
    nobullet(tf.paragraphs[0])


COUNTER = {"n": 1}


def new_slide(title=None, title_size=32):
    slide = prs.slides.add_slide(BLANK)
    if title is not None:
        tf = textbox(slide, 1.02, 0.72, 11.4, 0.75)
        p = tf.paragraphs[0]
        nobullet(p)
        r = p.add_run()
        r.text = title
        style(r, title_size, bold=True)
    footer(slide, COUNTER["n"])
    COUNTER["n"] += 1
    return slide


def bullets(slide, items, top=1.85, left=0.95, width=11.6, size=18, spacing=10):
    """items: list of (text, level) or plain strings (level 0)."""
    tf = textbox(slide, left, top, width, 7.0 - top - 0.5)
    first = True
    for item in items:
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bullet(p, level)
        p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = text
        style(r, size - (2 if level else 0))
    return tf


def _picture_centred(slide, path, centre_x, top, height):
    """Place a portrait screenshot centred on a column, not left-aligned to it."""
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    width = height * w / float(h)
    slide.shapes.add_picture(path, Inches(centre_x - width / 2.0),
                             Inches(top), height=Inches(height))


# ----------------------------------------------------------------------
# 1. Title
# ----------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
slide.shapes.add_picture(LOGO, Inches(4.22), Inches(0.42), width=Inches(4.90))

tf = textbox(slide, 2.0, 2.05, 9.3, 0.9)
for i, line in enumerate(["B.Tech Computer Science and Engineering",
                          "BCSE497J \u2013 Project-I"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    nobullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = line
    style(r, 20, bold=True)

tf = textbox(slide, 1.2, 3.02, 10.9, 1.1)
p = tf.paragraphs[0]
nobullet(p)
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = ("System and Method for Selective Multimodal Object-Detection "
          "Guidance for Visually Impaired Users")
style(r, 28, bold=True)

tf = textbox(slide, 1.0, 4.35, 8.5, 1.6)
p = tf.paragraphs[0]
nobullet(p)
r = p.add_run()
r.text = "Team members:"
style(r, 20)

MEMBERS = [("Aditya Sridhar", "23BCE2269"),
           ("Goureesankar S Nair", "23BCE2217"),
           ("Yuvan Mathan", "23BCE2106")]
for name, _reg in MEMBERS:
    p = tf.add_paragraph()
    bullet(p)
    r = p.add_run()
    r.text = name
    style(r, 20, bold=True)

# The registration numbers sit in their own box so they line up as a column;
# tab stops inside a bulleted list do not align reliably across renderers.
tfr = textbox(slide, 4.62, 4.675, 2.6, 1.2)
for i, (_name, reg) in enumerate(MEMBERS):
    p = tfr.paragraphs[0] if i == 0 else tfr.add_paragraph()
    nobullet(p)
    r = p.add_run()
    r.text = "(" + reg + ")"
    style(r, 20, bold=True)

tf = textbox(slide, 1.0, 6.05, 8.0, 0.8)
p = tf.paragraphs[0]
nobullet(p)
r = p.add_run()
r.text = "Faculty guide :"
style(r, 20)
p = tf.add_paragraph()
nobullet(p)
r = p.add_run()
r.text = "Dr. Vani Rajasekar"
style(r, 20, bold=True)

COUNTER["n"] = 2

# ----------------------------------------------------------------------
# 2. Aim
# ----------------------------------------------------------------------
slide = new_slide("Aim")
bullets(slide, [
    "To build an offline-first mobile assistant that uses a phone camera to tell a "
    "visually impaired user what is around them, where it is and how near it is, "
    "in short spoken phrases backed by stereo audio and haptic cues.",
    "To make the system say less rather than more: every layer \u2014 perception, "
    "guidance, transport and dialogue \u2014 abstains when it cannot verify what it "
    "would say, because a confident wrong sentence is unverifiable by a user who "
    "cannot look at the screen.",
    "To operate without any cloud service, so the user's home is never streamed to "
    "a third party and the assistant keeps working where there is no internet.",
    "To deliver a working Android application rather than a laboratory "
    "demonstration: a minimal non-visual interface driven by voice and gestures, "
    "answering in speech, stereo sonar and vibration, because the intended user "
    "cannot read a screen.",
], top=1.90, size=19, spacing=15)

# ----------------------------------------------------------------------
# 3. Abstract
# ----------------------------------------------------------------------
slide = new_slide("Abstract")
bullets(slide, [
    "Visually impaired users are poorly served by existing camera assistants because "
    "those assistants describe everything they see. A blind user cannot check a "
    "description against the room, so an over-full and occasionally wrong channel "
    "produces mistrust and is abandoned.",
    "BlindAssist is an Android application backed by a laptop inference server. It "
    "detects indoor objects with YOLOv8s plus a purpose-trained door/dustbin "
    "detector, converts each box into a direction, a clock bearing, a proximity "
    "bucket and a rough distance, and speaks only the single most relevant hazard.",
    "Beyond obstacle warnings it provides find, describe, count, clear-path, "
    "colour, light, read-text, page-summary, photo and object-memory capabilities, "
    "driven by offline speech recognition and, for unrehearsed phrasing, a local "
    "language-model router that may choose a capability but may never author a "
    "guidance sentence.",
    "The split between handset and laptop is an engineering result rather than a "
    "convenience: on-device inference measured 2.5 s per frame on the test phone "
    "against a 1 s guidance budget, so frames are JPEG-compressed and sent over a "
    "phone hotspot to a GPU server, while speech, sonar, haptics, text reading and "
    "the entire command parser stay on the phone and keep working if the link "
    "drops.",
    "Measured on recorded and live walks: 31/31 spoken directions correct with 0 "
    "phantom announcements, 21 ms two-model GPU inference, 26 KB JPEG frame "
    "transport, and a labelled naming index that corrected 104 of 105 COCO naming "
    "errors while renaming nothing in scenes it had not been taught.",
], top=1.68, size=15, spacing=10)

# ----------------------------------------------------------------------
# 4-5. Literature review
# ----------------------------------------------------------------------
LIT_HDR = ["Paper", "Objective", "Methodology", "Pros / Cons", "Findings"]
COL_W = [2.45, 2.05, 2.35, 2.40, 2.45]

LIT_1 = [
    ["Redmon et al., \u201cYou Only Look Once\u201d, CVPR 2016",
     "Real-time object detection in one network pass",
     "Single CNN regresses boxes and class scores over a grid; trained end-to-end",
     "+ Fast enough for live video\n\u2013 Coarse localisation; fixed closed vocabulary",
     "Established the detector family (now YOLOv8) this project uses for the "
     "\u201cwhere\u201d of every object"],
    ["Bigham et al., \u201cVizWiz\u201d, UIST 2010",
     "Answer blind users' visual questions in near real time",
     "Phone photo plus spoken question routed to paid remote sighted workers",
     "+ Accurate, open-ended answers\n\u2013 Latency, cost and a stranger sees the "
     "user's home",
     "Blind users ask about identity and text far more than geometry; privacy is a "
     "first-order design constraint"],
    ["Gurari et al., \u201cVizWiz Grand Challenge\u201d, CVPR 2018",
     "Benchmark VQA on photographs actually taken by blind users",
     "31,000 real user images with crowd answers, including unanswerable ones",
     "+ Realistic blur, framing and lighting\n\u2013 Models remain weak on the "
     "unanswerable class",
     "Automatic answering degrades sharply on blind-captured images; knowing when "
     "to decline is a distinct skill"],
    ["MacLeod et al., \u201cComputer-Generated Captions\u201d, CHI 2017",
     "Study how blind users treat automatically generated captions",
     "Controlled study varying caption phrasing and stated confidence",
     "+ Isolates trust from accuracy\n\u2013 Small sample, captions not "
     "interactive",
     "Users believe a confidently phrased caption even when it is wrong \u2014 the "
     "direct motivation for abstaining instead of guessing"],
    ["Lin et al., \u201cMicrosoft COCO\u201d, ECCV 2014",
     "Provide a large benchmark of everyday objects in context",
     "328,000 images exhaustively labelled over a fixed vocabulary of 80 classes",
     "+ The standard training set for detectors\n\u2013 80 closed classes, with no "
     "door, wardrobe, dustbin or laundry basket",
     "A detector must emit one of 80 words, so unlisted furniture is forced to the "
     "nearest wrong label \u2014 the naming problem this project sets out to fix"],
]

LIT_2 = [
    ["Sato et al., \u201cNavCog3\u201d, ASSETS 2017",
     "Turn-by-turn indoor navigation for blind travellers",
     "BLE beacon localisation plus a semantic map, spoken instructions",
     "+ Metre-level accuracy indoors\n\u2013 Requires beacons installed and a "
     "surveyed map",
     "Infrastructure-dependent guidance does not generalise to an arbitrary home; "
     "motivates a camera-only design"],
    ["Guerreiro et al., \u201cCaBot\u201d, ASSETS 2019",
     "Autonomous suitcase-shaped robot guiding a blind user",
     "LiDAR navigation with haptic handle and speech feedback",
     "+ Strong obstacle avoidance\n\u2013 Bulky, expensive, not a personal daily "
     "device",
     "Confirms multimodal (speech + haptic) output is preferred over speech alone "
     "while walking"],
    ["Meijer, \u201cThe vOICe\u201d, IEEE Trans. Biomed. Eng., 1992",
     "Represent an entire image as sound for blind users",
     "Left-to-right sweep mapping image columns to stereo position, rows to pitch",
     "+ Continuous, hands-free\n\u2013 Long training period; high cognitive load",
     "Stereo panning and pitch are learnable spatial carriers \u2014 the basis of "
     "this project's sonar mode"],
    ["Wickens & Dixon, Theor. Issues in Ergonomics Sci., 2007",
     "Quantify the operator cost of imperfect automated alerts",
     "Meta-analysis of automation reliability against human performance",
     "+ Gives a reliability threshold\n\u2013 Not specific to assistive devices",
     "Below roughly 70% reliability an alerting aid is worse than none; false "
     "alarms cost more than misses"],
    ["Schick et al., “Toolformer”, NeurIPS 2023",
     "Let a language model decide for itself when to call an external tool",
     "Self-supervised insertion of API calls, with the tool returning the value "
     "the model then uses",
     "+ Keeps facts outside the model\n– The model still writes the sentence "
     "around the returned value",
     "Supports routing speech to a capability, but the phrasing must stay outside "
     "the model — the basis of this project's authority boundary"],
]


def lit_table(slide, rows):
    n = len(rows) + 1
    tbl_shape = slide.shapes.add_table(
        n, 5, Inches(0.80), Inches(1.63), Inches(sum(COL_W)), Inches(0.84 * n))
    tbl = tbl_shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, w in enumerate(COL_W):
        tbl.columns[i].width = Inches(w)

    for c, head in enumerate(LIT_HDR):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TBL_HDR
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        p = cell.text_frame.paragraphs[0]
        nobullet(p)
        r = p.add_run()
        r.text = head
        style(r, 16, bold=True, color=WHITE)

    for ri, row in enumerate(rows, start=1):
        shade = TBL_A if ri % 2 == 1 else TBL_B
        for ci, text in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = shade
            cell.vertical_anchor = MSO_ANCHOR.TOP
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            tf = cell.text_frame
            tf.word_wrap = True
            for li, line in enumerate(text.split("\n")):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                nobullet(p)
                r = p.add_run()
                r.text = line
                style(r, 11, bold=(ci == 0))


slide = new_slide("Literature Review")
tf = textbox(slide, 0.82, 1.30, 11.5, 0.3)
nobullet(tf.paragraphs[0])
r = tf.paragraphs[0].add_run()
r.text = "Detection, remote assistance and the trust problem (1 of 2)"
style(r, 14, italic=True)
lit_table(slide, LIT_1)

slide = new_slide("Literature Review")
tf = textbox(slide, 0.82, 1.30, 11.5, 0.3)
nobullet(tf.paragraphs[0])
r = tf.paragraphs[0].add_run()
r.text = "Navigation aids, non-visual output and automation reliability (2 of 2)"
style(r, 14, italic=True)
lit_table(slide, LIT_2)

# ----------------------------------------------------------------------
# 6. Research gap
# ----------------------------------------------------------------------
slide = new_slide("Research Gap")
bullets(slide, [
    "Existing assistants optimise coverage, not silence. Seeing AI, Envision and "
    "Lookout narrate what they see; none has a principled rule for declining to "
    "speak, so the user cannot tell a confident sentence from a correct one.",
    "Verification is asymmetric and unaddressed. A sighted user glances at the "
    "screen to reject a wrong label; a blind user cannot, so the cost of a false "
    "statement is far higher than the literature's accuracy metrics capture.",
    "Confidence is treated as a proxy for correctness, and it is not. Measured on "
    "our own clips, a wrongly named dustbin scored 0.94 while a correctly named "
    "chair scored 0.92 \u2014 the bands overlap, so no confidence threshold "
    "separates right names from wrong ones.",
    "Closed detector vocabularies force wrong words. COCO has no class for a "
    "wardrobe, dustbin or laundry basket, so the network is compelled to emit the "
    "nearest of 80 words; a laundry basket becomes \u201chandbag\u201d and is then "
    "silently discarded by the pipeline.",
    "Indoor navigation aids need infrastructure. Beacon and map-based systems do "
    "not transfer to an arbitrary home, and GPS carries 10\u201350 m of error "
    "indoors against a 10\u201315 m flat.",
    "Language-model assistants are ungrounded. Left free to phrase perceptual "
    "claims, a small local model fabricated content in 42.5% of our free-text "
    "ablation cases \u2014 unacceptable when the claim concerns a staircase.",
], top=1.8, size=15, spacing=11)

# ----------------------------------------------------------------------
# 7. Objectives
# ----------------------------------------------------------------------
slide = new_slide("Objectives")
bullets(slide, [
    "Detect the indoor objects that matter for mobility and search, including "
    "classes absent from COCO, by combining a general detector with a "
    "purpose-trained door/dustbin model.",
    "Convert every detection into actionable geometry \u2014 3\u00d73 zone, "
    "clock-face bearing derived from the true 65\u00b0 camera field of view, "
    "proximity bucket and monocular distance estimate.",
    "Announce at most one object per moment, with anti-repetition, escalation "
    "override and a focus rule that prevents a routine warning cutting across a "
    "task the user deliberately asked for.",
    "Provide non-visual output beyond speech: stereo-panned sonar beeps whose rate "
    "and pitch track proximity, and pulse-count haptics encoding side.",
    "Accept unrehearsed natural speech through a two-tier router in which a local "
    "language model may choose a capability but may never author a perceptual "
    "claim or a guidance instruction.",
    "Keep the entire pipeline offline and quantify every claim on recorded clips, "
    "replayable field walks and a 200-utterance frozen evaluation set.",
], top=1.75, size=15, spacing=9)

tf = textbox(slide, 0.95, 5.80, 11.6, 1.0)
lines = [
    ("Identified Sustainable Development Goals:  ",
     "SDG 3 \u2013 Good Health and Well-being  |  SDG 10 \u2013 Reduced "
     "Inequalities  |  SDG 11 \u2013 Sustainable Cities and Communities"),
    ("Technology Readiness Level:  ",
     "TRL 5 \u2013 technology validated in a relevant environment (working "
     "handset build, field-walked in real rooms)"),
    ("Expected outcomes:  ",
     "Patent filing (\u201cselective multimodal guidance\u201d) and a Scopus-indexed "
     "conference / journal paper"),
]
for i, (lead, rest) in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    nobullet(p)
    p.space_after = Pt(3)
    r = p.add_run()
    r.text = lead
    style(r, 13, bold=True)
    r = p.add_run()
    r.text = rest
    style(r, 13)

# ----------------------------------------------------------------------
# 8. Architecture
# ----------------------------------------------------------------------
slide = new_slide("Framework / Architecture / Block Diagram")


def box(slide, left, top, width, height, text, size=10, bold=False, dashed=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left),
                                Inches(top), Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = BLACK
    sh.line.width = Pt(1.0)
    sh.adjustments[0] = 0.08
    if dashed:
        sh.line.dash_style = 4  # msoLineDash
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nobullet(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        style(r, size, bold=(bold and i == 0))
    return sh


def arrow(slide, x1, y1, x2, y2):
    con = slide.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    con.line.color.rgb = BLACK
    con.line.width = Pt(1.25)
    ln = con.line._get_or_add_ln()
    tail = ln.find(qn("a:tailEnd"))
    if tail is None:
        tail = ln.makeelement(qn("a:tailEnd"), {})
        ln.append(tail)
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return con


def label(slide, left, top, width, text, size=10, italic=False, align=PP_ALIGN.CENTER):
    tf = textbox(slide, left, top, width, 0.3)
    p = tf.paragraphs[0]
    nobullet(p)
    p.alignment = align
    r = p.add_run()
    r.text = text
    style(r, size, italic=italic)


# --- Row A: perception path (phone -> server -> phone)
label(slide, 0.55, 1.50, 4.0, "A.  Perception path", 13, align=PP_ALIGN.LEFT)
ya, ha = 1.85, 0.72
xs = [0.55, 2.62, 4.69, 6.76, 8.83, 10.90]
wb = 1.85
box(slide, xs[0], ya, wb, ha, "Phone camera\nYUV420 stream, 720p")
box(slide, xs[1], ya, wb, ha, "Hardware JPEG\n506 KB \u2192 26 KB")
box(slide, xs[2], ya, wb, ha, "Wi-Fi POST /infer\nUDP auto-discovery")
box(slide, xs[3], ya, wb, ha, "Detector ensemble\nYOLOv8s + door/dustbin")
box(slide, xs[4], ya, wb, ha, "Cross-model merge\nIoU, calibrated margin")
box(slide, xs[5], ya, wb, ha, "Naming head\nembedding nearest-neighbour")
for i in range(5):
    arrow(slide, xs[i] + wb, ya + ha / 2, xs[i + 1], ya + ha / 2)

# --- Row B: on-phone reasoning and output
yb = 3.05
box(slide, xs[5], yb, wb, ha, "Position analysis\nzone, clock, proximity, distance")
box(slide, xs[4], yb, wb, ha, "Guidance engine\nWalk / Find, anti-spam")
box(slide, xs[3], yb, wb, ha, "Speech policy\npriority + focus arbitration")
box(slide, xs[2], yb, wb, ha, "Text-to-speech\nshort phrases")
box(slide, xs[1], yb, wb, ha, "Sonar audio\nstereo pan, rate, pitch")
box(slide, xs[0], yb, wb, ha, "Haptics\npulse count = side")
arrow(slide, xs[5] + wb / 2, ya + ha, xs[5] + wb / 2, yb)
for i in range(5, 0, -1):
    arrow(slide, xs[i], yb + ha / 2, xs[i - 1] + wb, yb + ha / 2)

# --- Row C: dialogue path
label(slide, 0.55, 4.10, 4.0, "B.  Dialogue path", 13, align=PP_ALIGN.LEFT)
yc = 4.45
box(slide, xs[0], yc, wb, ha, "Microphone\nalways-on")
box(slide, xs[1], yc, wb, ha, "Vosk offline ASR\ngrammar-constrained")
box(slide, xs[2], yc, wb, ha, "Tier 0 parser\nexact phrases, 5 \u00b5s")
box(slide, xs[3], yc, wb, ha, "Tier 1 router\nllama3.2:1b, on miss")
box(slide, xs[4], yc, wb, ha, "Authority boundary\ntool + argument validation")
box(slide, xs[5], yc, wb, ha, "Capability executed\nor abstain")
for i in range(5):
    arrow(slide, xs[i] + wb, yc + ha / 2, xs[i + 1], yc + ha / 2)
arrow(slide, xs[5] + wb / 2, yc, xs[5] + wb / 2, yb + ha)

# --- Row D: on-demand capture path
label(slide, 0.55, 5.28, 4.5, "C.  On-demand capabilities", 13, align=PP_ALIGN.LEFT)
yd = 5.72
box(slide, xs[0], yd, wb, ha, "Still capture\nstream paused, autofocus")
box(slide, xs[1], yd, wb, ha, "ML Kit OCR\non-device, offline")
box(slide, xs[2], yd, wb, ha, "POST /summarise\nfigures checked")
box(slide, xs[3], yd, wb, ha, "Colour and light\ncentre patch, luma")
box(slide, xs[4], yd, wb, ha, "Object memory\nday-long, persisted")
box(slide, xs[5], yd, wb, ha, "Photo to gallery\nfor a sighted helper")
arrow(slide, xs[0] + wb, yd + ha / 2, xs[1], yd + ha / 2)
arrow(slide, xs[1] + wb, yd + ha / 2, xs[2], yd + ha / 2)

label(slide, 0.55, 6.55, 12.2, "Every block runs offline. The dialogue tier may "
      "select a capability; it never authors a guidance sentence, and any "
      "unverifiable output is replaced by silence.", 12, italic=True,
      align=PP_ALIGN.LEFT)

# ----------------------------------------------------------------------
# 9. Functional requirements
# ----------------------------------------------------------------------
slide = new_slide("Functional Requirements")
bullets(slide, [
    "FR1 \u2014 Walk mode: the system shall continuously detect obstacles and "
    "announce at most one per moment, chosen by proximity, then centrality, then "
    "size, and only when close or very close.",
    "FR2 \u2014 Find mode: on a spoken class name the system shall report the "
    "target's bearing, proximity and rough distance on the first frame it is seen, "
    "then return to walk mode.",
    "FR3 \u2014 Scene understanding: the system shall provide describe, count "
    "and clear-path answers grouped by direction on demand.",
    "FR4 \u2014 Non-visual output: the system shall provide stereo sonar beeps "
    "whose tick rate and pitch rise with proximity, and haptic pulses whose count "
    "encodes left, ahead or right.",
    "FR5 \u2014 Voice control: all capabilities shall be reachable hands-free "
    "through offline speech recognition, with a trigger word opening a free-speech "
    "window for unrehearsed phrasing.",
    "FR6 \u2014 Reading: the system shall read printed text aloud from a still "
    "capture and, on request, summarise a page with every figure verified against "
    "the source.",
    "FR7 \u2014 Memory: the system shall recall where a named object was last "
    "seen, stating the age of that observation before the location.",
    "FR8 \u2014 Fail-safe: on loss of the inference link the system shall pause "
    "guidance and say so, rather than treating no data as an empty room.",
    "FR9 \u2014 Abstention: the system shall remain silent whenever a name, a "
    "distance or a routed request cannot be verified, and shall never speak a "
    "perceptual claim originating in the language model.",
], top=1.75, size=14, spacing=8)

# ----------------------------------------------------------------------
# 10. Modules
# ----------------------------------------------------------------------
slide = new_slide("Modules")
bullets(slide, [
    "M1 Detection \u2014 YOLOv8s over COCO plus a self-trained YOLOv8n for door "
    "and dustbin; per-class thresholds and per-model non-maximum suppression "
    "(infer_server.py, detector.dart).",
    "M2 Cross-model merge \u2014 removes the same object detected twice under two "
    "names, comparing confidence as margin above each model's own floor "
    "(detect_merge.py).",
    "M3 Naming head \u2014 re-decides the word for a crop from an embedding matched "
    "against 280 user-labelled examples, with four abstention rules "
    "(name_index.py, build_name_index.py).",
    "M4 Position analysis \u2014 pure logic converting a box into zone, clock "
    "bearing, proximity bucket and monocular distance (position.py / position.dart).",
    "M5 Guidance engine \u2014 walk and find decision logic, persistence, "
    "cooldowns, escalation, scene summary, clear path, object memory "
    "(decision.py / decision.dart).",
    "M6 Speech policy \u2014 four priorities plus a focus hold, so a task the user "
    "asked for owns the channel until it finishes (speech_policy.py / .dart).",
    "M7 Voice and agent \u2014 Vosk grammar recognition, tier-0 parser and the "
    "tier-1 router behind a validating authority boundary (voice.py, agent.py, "
    "agent_server.py).",
    "M8 Output \u2014 text-to-speech with stop-before-speak, WebAudio/native sonar "
    "and vibration (speaker.dart, sonar.dart).",
    "M9 On-demand capture \u2014 OCR, page summary, colour and light naming, photo "
    "to gallery (ocr.dart, text_summary.py, colour_naming.py).",
    "M10 Harnesses \u2014 Flask web UI for development, replayable clip evaluation, "
    "445 Python and 271 Dart unit tests (webapp.py, verify_namer.py, eval_agent.py).",
], top=1.72, size=13, spacing=7)

# ----------------------------------------------------------------------
# 11. Experiments and results (quantitative)
# ----------------------------------------------------------------------
slide = new_slide("Experiments and Results")

RES = [
    ["Experiment", "Setup", "Result"],
    ["Spoken direction accuracy",
     "Recorded room clips, every announcement keyframe reviewed by hand",
     "31/31 directions correct; 0/31 phantom announcements; 6/31 wrong object name "
     "(COCO vocabulary)"],
    ["Detector latency",
     "Both models, 640 px, RTX 3050 laptop GPU, 12 frames, median",
     "21.2 ms per frame combined (47 FPS) against 256.5 ms on CPU \u2014 a 12\u00d7 "
     "reduction"],
    ["Frame transport",
     "Handset to server over a phone hotspot",
     "Hardware JPEG cut 506 KB to 26 KB; upload 320\u2013510 ms of timeouts removed, "
     "283/283 frames delivered"],
    ["Embedding naming head",
     "280 labelled crops, 17 labels; leave-one-out and 8-clip replay at stride 1",
     "49/49 leave-one-out names correct at margin 0.15; 105 renames on clips, 104 "
     "correct; 0 renames in unseen scenes"],
    ["Find-mode responsiveness",
     "User's room replayed at the handset's true 2 FPS",
     "First correct answer moved from 19.5 s to the first frame; 12/12 classes "
     "answered, 0 false \u201cnot visible\u201d"],
    ["Command routing (n = 200)",
     "Frozen evaluation set; keyword baseline against the two-tier router",
     "Overall 39.5% \u2192 53.0%; paraphrase 0% \u2192 47.1%; canonical commands "
     "held at 100%; tier 0 routes in 5 \u00b5s"],
    ["Authority boundary",
     "Free-text ablation letting the model phrase perceptual claims",
     "42.5% of replies fabricated content; with the boundary enforced, 0 guidance "
     "strings originated in the model"],
]

tbl = slide.shapes.add_table(len(RES), 3, Inches(0.92), Inches(1.55),
                             Inches(11.5), Inches(0.5 * len(RES))).table
for i, w in enumerate([2.55, 3.85, 5.10]):
    tbl.columns[i].width = Inches(w)
for ri, row in enumerate(RES):
    for ci, text in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = (TBL_HDR if ri == 0
                                    else (TBL_A if ri % 2 == 1 else TBL_B))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.07)
        cell.margin_right = Inches(0.07)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        p = cell.text_frame.paragraphs[0]
        nobullet(p)
        r = p.add_run()
        r.text = text
        style(r, 11 if ri else 14, bold=(ri == 0 or ci == 0),
              color=(WHITE if ri == 0 else BLACK))

# ----------------------------------------------------------------------
# 12. Experiments and results - walk mode screenshots
# ----------------------------------------------------------------------
slide = new_slide("Experiments and Results")
label(slide, 0.95, 1.50, 11.6, "Walk mode on the handset: one hazard announced "
      "per moment, with a clock bearing and a sidestep when very close.",
      15, italic=True, align=PP_ALIGN.LEFT)

shots = [
    (os.path.join(IMG, "walk1.png"),
     "Cluttered room. Four objects detected; only the nearest is spoken:\n"
     "\u201cObstacle very close at 12 o'clock, move slightly right\u201d.\n"
     "The name is withheld because the detector's word is not trusted."),
    (os.path.join(IMG, "walk2.png"),
     "Desk scene. The self-trained model contributes the dustbin, which COCO "
     "cannot name:\n\u201cDustbin at 12 o'clock, close\u201d."),
    (os.path.join(IMG, "walk3.png"),
     "Floor-level hazard at the user's feet:\n\u201cSuitcase very close at 12 "
     "o'clock, move slightly right\u201d, escalated on approach."),
]
COLS = [(0.53, 4.15), (4.68, 4.15), (8.83, 4.15)]
for (col_left, col_w), (path, caption) in zip(COLS, shots):
    _picture_centred(slide, path, col_left + col_w / 2.0, 2.05, 3.55)
    tfc = textbox(slide, col_left, 5.72, col_w, 1.2)
    for li, line in enumerate(caption.split("\n")):
        p = tfc.paragraphs[0] if li == 0 else tfc.add_paragraph()
        nobullet(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        style(r, 11)

# ----------------------------------------------------------------------
# 13. Experiments and results - other capabilities
# ----------------------------------------------------------------------
slide = new_slide("Experiments and Results")
label(slide, 0.95, 1.50, 11.6, "On-demand capabilities: reading printed text and "
      "naming a colour, both captured on the handset.", 15, italic=True,
      align=PP_ALIGN.LEFT)

_picture_centred(slide, os.path.join(IMG, "ocr.png"), 3.825, 2.05, 3.55)
_picture_centred(slide, os.path.join(IMG, "colour.png"), 9.525, 2.05, 3.55)

for left, caption in [
    (0.95, "Read: the stream is paused, a still is captured at 720p with "
           "autofocus, and on-device OCR speaks the text. Recognition is "
           "reported verbatim \u2014 nothing is invented to smooth it over."),
    (6.65, "Colour: the centre patch is averaged and named only when the "
           "exposure allows it. \u201cBlue\u201d is spoken here; under poor "
           "light the capability abstains rather than guess a shade."),
]:
    tfc = textbox(slide, left, 5.75, 5.75, 1.1)
    p = tfc.paragraphs[0]
    nobullet(p)
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = caption
    style(r, 12)

# ----------------------------------------------------------------------
# 14. Conclusion
# ----------------------------------------------------------------------
slide = new_slide("Conclusion")
bullets(slide, [
    "A complete offline assistive pipeline has been built and field-walked: phone "
    "camera to JPEG transport to a two-model detector with an embedding naming "
    "head, then position analysis, guidance arbitration and speech, sonar and "
    "haptic output, all reachable by voice.",
    "Selective abstention is the contribution and it is measurable. Withholding a "
    "name the system cannot verify, pausing guidance when the link drops, "
    "declining a distance from a clipped box and refusing a routed action whose "
    "argument the user never said each removed a class of confident error.",
    "Confidence was shown not to be a correctness signal on our own data \u2014 a "
    "misnamed dustbin scored 0.94 against a correct chair at 0.92 \u2014 and was "
    "replaced by a calibrated embedding margin that renamed 104 of 105 objects "
    "correctly and stayed silent in rooms it had not been taught.",
    "Tiering beats a language model alone: canonical commands stay at 100% and "
    "route in microseconds, while paraphrase accuracy rises from 0% to 47%, and "
    "the authority boundary keeps every spoken guidance string deterministic.",
    "Honest limitations. Over-triggering on out-of-scope speech remains at 55%, "
    "the naming index is room-specific until crops from that room are labelled, "
    "clock bearings cannot be finer than three zones at a 65\u00b0 field of view, "
    "and no blind user has yet tested the system.",
    "Next: label the harvested crops from the newest room, run the field-test "
    "protocol including the server-kill drill, measure battery and thermal cost, "
    "and recruit visually impaired participants for the evaluation the design was "
    "argued from.",
], top=1.78, size=14, spacing=10)

# ----------------------------------------------------------------------
# 15-16. References
# ----------------------------------------------------------------------
REFS = [
    "J. Redmon, S. Divvala, R. Girshick, and A. Farhadi, \u201cYou only look once: "
    "Unified, real-time object detection,\u201d in Proc. IEEE Conf. Computer Vision "
    "and Pattern Recognition (CVPR), Las Vegas, NV, USA, 2016, pp. 779\u2013788.",

    "G. Jocher, A. Chaurasia, and J. Qiu, \u201cUltralytics YOLO,\u201d software, "
    "version 8, 2023. [Online]. Available: https://github.com/ultralytics/ultralytics",

    "T.-Y. Lin, M. Maire, S. Belongie, J. Hays, P. Perona, D. Ramanan, P. Doll\u00e1r, "
    "and C. L. Zitnick, \u201cMicrosoft COCO: Common objects in context,\u201d in "
    "Proc. European Conf. Computer Vision (ECCV), Zurich, Switzerland, 2014, "
    "pp. 740\u2013755.",

    "J. P. Bigham, C. Jayant, H. Ji, G. Little, A. Miller, R. C. Miller, R. Miller, "
    "A. Tatarowicz, B. White, S. White, and T. Yeh, \u201cVizWiz: Nearly real-time "
    "answers to visual questions,\u201d in Proc. 23rd ACM Symp. User Interface "
    "Software and Technology (UIST), New York, NY, USA, 2010, pp. 333\u2013342.",

    "D. Gurari, Q. Li, A. J. Stangl, A. Guo, C. Lin, K. Grauman, J. Luo, and "
    "J. P. Bigham, \u201cVizWiz grand challenge: Answering visual questions from "
    "blind people,\u201d in Proc. IEEE/CVF Conf. Computer Vision and Pattern "
    "Recognition (CVPR), Salt Lake City, UT, USA, 2018, pp. 3608\u20133617.",

    "H. MacLeod, C. L. Bennett, M. R. Morris, and E. Cutrell, \u201cUnderstanding "
    "blind people's experiences with computer-generated captions of social media "
    "images,\u201d in Proc. CHI Conf. Human Factors in Computing Systems, Denver, "
    "CO, USA, 2017, pp. 5988\u20135999.",

    "A. Stangl, M. R. Morris, and D. Gurari, \u201c\u2018Person, shoes, tree. Is "
    "the person naked?\u2019 What people with vision impairments want in image "
    "descriptions,\u201d in Proc. CHI Conf. Human Factors in Computing Systems, "
    "Honolulu, HI, USA, 2020, pp. 1\u201313.",

    "D. Ahmetovic, C. Gleason, C. Ruan, K. Kitani, H. Takagi, and C. Asakawa, "
    "\u201cNavCog: A navigational cognitive assistant for the blind,\u201d in Proc. "
    "18th Int. Conf. Human-Computer Interaction with Mobile Devices and Services "
    "(MobileHCI), Florence, Italy, 2016, pp. 90\u201399.",

    "D. Sato, U. Oh, K. Naito, H. Takagi, K. Kitani, and C. Asakawa, "
    "\u201cNavCog3: An evaluation of a smartphone-based blind indoor navigation "
    "assistant with semantic features in a large-scale environment,\u201d in Proc. "
    "19th Int. ACM SIGACCESS Conf. Computers and Accessibility (ASSETS), Baltimore, "
    "MD, USA, 2017, pp. 270\u2013279.",

    "J. Guerreiro, D. Sato, S. Asakawa, H. Dong, K. M. Kitani, and C. Asakawa, "
    "\u201cCaBot: Designing and evaluating an autonomous navigation robot for blind "
    "people,\u201d in Proc. 21st Int. ACM SIGACCESS Conf. Computers and "
    "Accessibility (ASSETS), Pittsburgh, PA, USA, 2019, pp. 68\u201382.",

    "B. Kuriakose, R. Shrestha, and F. E. Sandnes, \u201cTools and technologies for "
    "blind and visually impaired navigation support: A review,\u201d IETE Technical "
    "Review, vol. 39, no. 1, pp. 3\u201318, 2022.",

    "P. B. L. Meijer, \u201cAn experimental system for auditory image "
    "representations,\u201d IEEE Trans. Biomedical Engineering, vol. 39, no. 2, "
    "pp. 112\u2013121, Feb. 1992.",

    "J. B. F. van Erp, H. A. H. C. van Veen, C. Jansen, and T. Dobbins, "
    "\u201cWaypoint navigation with a vibrotactile waist belt,\u201d ACM Trans. "
    "Applied Perception, vol. 2, no. 2, pp. 106\u2013117, Apr. 2005.",

    "J. D. Lee and K. A. See, \u201cTrust in automation: Designing for appropriate "
    "reliance,\u201d Human Factors, vol. 46, no. 1, pp. 50\u201380, 2004.",

    "C. D. Wickens and S. R. Dixon, \u201cThe benefits of imperfect diagnostic "
    "automation: A synthesis of the literature,\u201d Theoretical Issues in "
    "Ergonomics Science, vol. 8, no. 3, pp. 201\u2013212, 2007.",

    "C. K. Chow, \u201cOn optimum recognition error and reject tradeoff,\u201d IEEE "
    "Trans. Information Theory, vol. 16, no. 1, pp. 41\u201346, Jan. 1970.",

    "Y. Geifman and R. El-Yaniv, \u201cSelective classification for deep neural "
    "networks,\u201d in Advances in Neural Information Processing Systems (NeurIPS), "
    "Long Beach, CA, USA, 2017, pp. 4878\u20134887.",

    "Z. Ji, N. Lee, R. Frieske, T. Yu, D. Su, Y. Xu, E. Ishii, Y. Bang, A. Madotto, "
    "and P. Fung, \u201cSurvey of hallucination in natural language generation,\u201d "
    "ACM Computing Surveys, vol. 55, no. 12, pp. 1\u201338, 2023.",

    "T. Schick, J. Dwivedi-Yu, R. Dess\u00ec, R. Raileanu, M. Lomeli, L. Zettlemoyer, "
    "N. Cancedda, and T. Scialom, \u201cToolformer: Language models can teach "
    "themselves to use tools,\u201d in Advances in Neural Information Processing "
    "Systems (NeurIPS), New Orleans, LA, USA, 2023.",

    "D. Povey, A. Ghoshal, G. Boulianne, L. Burget, O. Glembek, N. Goel, "
    "M. Hannemann, P. Motl\u00ed\u010dek, Y. Qian, P. Schwarz, J. Silovsk\u00fd, "
    "G. Stemmer, and K. Vesel\u00fd, \u201cThe Kaldi speech recognition "
    "toolkit,\u201d in Proc. IEEE Workshop Automatic Speech Recognition and "
    "Understanding (ASRU), Waikoloa, HI, USA, 2011.",
]


def ref_slide(title, items, start):
    slide = new_slide(title)
    tf = textbox(slide, 0.7, 1.55, 12.0, 5.3)
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nobullet(p)
        p.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(304800))
        pPr.set("indent", str(-304800))
        r = p.add_run()
        r.text = "[" + str(start + i) + "] " + text
        style(r, 12)


ref_slide("References", REFS[:10], 1)
ref_slide("References", REFS[10:], 11)

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
